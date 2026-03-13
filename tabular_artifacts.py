from __future__ import annotations

import csv
import io
import os
import tempfile
from collections.abc import Iterator

from config import UPLOAD_TABULAR_ARTIFACT_BATCH_ROWS
from tabular_loader import (
    TABULAR_PREVIEW_CHAR_LIMIT,
    TABULAR_PREVIEW_ROW_LIMIT,
    TabularLoaderError,
    get_tabular_extension,
    _normalize_header_row,
    _normalize_row,
    _preview_payload,
    _row_values_from_sequence,
    _temporary_tabular_file,
)


def build_tabular_artifact(
    raw_bytes: bytes,
    filename: str,
    *,
    batch_rows: int = UPLOAD_TABULAR_ARTIFACT_BATCH_ROWS,
) -> dict:
    import duckdb

    columns, row_iter = _iter_tabular_rows(raw_bytes, filename)
    if not columns:
        raise TabularLoaderError("Não foi possível criar artefacto tabular sem colunas.")

    safe_batch_rows = max(500, min(int(batch_rows or 0), 50_000))
    temp_handle = tempfile.NamedTemporaryFile(delete=False, suffix=".parquet")
    temp_handle.close()
    temp_path = temp_handle.name
    row_count = 0
    conn = duckdb.connect(database=":memory:")
    try:
        conn.execute(f"CREATE TABLE uploaded ({_duckdb_column_defs(columns)})")
        batch: list[list[str]] = []
        for row in row_iter:
            batch.append(_normalize_row(row, len(columns)))
            row_count += 1
            if len(batch) >= safe_batch_rows:
                _insert_batch(conn, columns, batch)
                batch = []
        if batch:
            _insert_batch(conn, columns, batch)

        conn.execute(
            "COPY uploaded TO ? (FORMAT PARQUET, COMPRESSION ZSTD)",
            [temp_path],
        )
        with open(temp_path, "rb") as fh:
            artifact_bytes = fh.read()
        return {
            "format": "parquet",
            "row_count": row_count,
            "columns": columns,
            "artifact_bytes": artifact_bytes,
        }
    finally:
        conn.close()
        try:
            os.unlink(temp_path)
        except Exception:
            pass


def load_tabular_artifact_dataset(
    artifact_bytes: bytes,
    *,
    max_rows: int,
) -> dict:
    import duckdb

    safe_limit = max(1, int(max_rows or 1))
    with _temporary_tabular_file(artifact_bytes, ".parquet") as temp_path:
        conn = duckdb.connect(database=":memory:")
        try:
            row_count = int(conn.execute("SELECT COUNT(*) FROM read_parquet(?)", [temp_path]).fetchone()[0] or 0)
            query = f"SELECT * FROM read_parquet(?) LIMIT {safe_limit}"
            cursor = conn.execute(query, [temp_path])
            rows = cursor.fetchall()
            columns = [str(col[0]) for col in (cursor.description or [])]
        finally:
            conn.close()

    records = [{column: _duckdb_value_to_string(row[idx]) for idx, column in enumerate(columns)} for row in rows]
    return {
        "columns": columns,
        "records": records,
        "row_count": row_count,
        "rows_loaded": len(records),
        "truncated": row_count > len(records),
        "delimiter": "\t",
    }


def load_tabular_artifact_preview(
    artifact_bytes: bytes,
    *,
    preview_rows: int = TABULAR_PREVIEW_ROW_LIMIT,
    preview_char_limit: int = TABULAR_PREVIEW_CHAR_LIMIT,
) -> dict:
    import duckdb

    safe_preview_rows = max(1, int(preview_rows or 1))
    with _temporary_tabular_file(artifact_bytes, ".parquet") as temp_path:
        conn = duckdb.connect(database=":memory:")
        try:
            row_count = int(conn.execute("SELECT COUNT(*) FROM read_parquet(?)", [temp_path]).fetchone()[0] or 0)
            query = f"SELECT * FROM read_parquet(?) LIMIT {safe_preview_rows}"
            cursor = conn.execute(query, [temp_path])
            rows = cursor.fetchall()
            columns = [str(col[0]) for col in (cursor.description or [])]
        finally:
            conn.close()

    sample_rows: list[list[str]] = []
    preview_lines = ["\t".join(columns)]
    truncated = False
    for row in rows:
        normalized = [_duckdb_value_to_string(value) for value in row]
        if len(sample_rows) < safe_preview_rows:
            sample_rows.append(normalized)
        line = "\t".join(normalized)
        current_size = sum(len(item) for item in preview_lines) + max(0, len(preview_lines) - 1)
        projected = current_size + 1 + len(line)
        if projected <= preview_char_limit:
            preview_lines.append(line)
        else:
            truncated = True
    return _preview_payload(columns, sample_rows, row_count, "\t", preview_lines, truncated)


def export_tabular_artifact_as_csv_bytes(artifact_bytes: bytes) -> bytes:
    import duckdb

    if not artifact_bytes:
        raise TabularLoaderError("Artefacto tabular vazio.")

    with _temporary_tabular_file(artifact_bytes, ".parquet") as parquet_path:
        with tempfile.NamedTemporaryFile(prefix="dbde_artifact_", suffix=".csv", delete=False) as tmp_csv:
            csv_path = tmp_csv.name
        conn = duckdb.connect(database=":memory:")
        try:
            safe_parquet_path = str(parquet_path).replace("'", "''")
            safe_csv_path = str(csv_path).replace("'", "''")
            conn.execute(
                f"COPY (SELECT * FROM read_parquet('{safe_parquet_path}')) "
                f"TO '{safe_csv_path}' (FORMAT CSV, HEADER, DELIMITER ',')",
            )
            with open(csv_path, "rb") as fh:
                return fh.read()
        finally:
            conn.close()
            try:
                os.unlink(csv_path)
            except OSError:
                pass


def iter_tabular_artifact_batches(
    artifact_bytes: bytes,
    *,
    columns: list[str] | None = None,
    batch_rows: int = 5000,
) -> Iterator[list[dict[str, str]]]:
    import duckdb

    if not artifact_bytes:
        raise TabularLoaderError("Artefacto tabular vazio.")

    safe_batch_rows = max(100, min(int(batch_rows or 0), 50_000))
    selected_columns = [str(col or "").strip() for col in (columns or []) if str(col or "").strip()]

    with _temporary_tabular_file(artifact_bytes, ".parquet") as temp_path:
        conn = duckdb.connect(database=":memory:")
        try:
            if selected_columns:
                select_clause = ", ".join(_duckdb_ident(column) for column in selected_columns)
            else:
                select_clause = "*"
            cursor = conn.execute(f"SELECT {select_clause} FROM read_parquet(?)", [temp_path])
            result_columns = [str(col[0]) for col in (cursor.description or [])]
            while True:
                rows = cursor.fetchmany(safe_batch_rows)
                if not rows:
                    break
                yield [
                    {
                        column: _duckdb_value_to_string(row[idx])
                        for idx, column in enumerate(result_columns)
                    }
                    for row in rows
                ]
        finally:
            conn.close()


def _insert_batch(conn, columns: list[str], batch: list[list[str]]) -> None:
    placeholders = ", ".join(["?"] * len(columns))
    conn.executemany(f"INSERT INTO uploaded VALUES ({placeholders})", batch)


def _duckdb_column_defs(columns: list[str]) -> str:
    return ", ".join(f"{_duckdb_ident(column)} VARCHAR" for column in columns)


def _duckdb_ident(value: str) -> str:
    return '"' + str(value or "").replace('"', '""') + '"'


def _duckdb_value_to_string(value) -> str:
    if value is None:
        return ""
    return str(value)


def _iter_tabular_rows(raw_bytes: bytes, filename: str) -> tuple[list[str], Iterator[list[str]]]:
    ext = get_tabular_extension(filename)
    if ext == ".csv":
        return _iter_delimited_rows(raw_bytes, delimiter_hint=None)
    if ext == ".tsv":
        return _iter_delimited_rows(raw_bytes, delimiter_hint="\t")
    if ext == ".xlsx":
        return _iter_xlsx_rows(raw_bytes)
    if ext == ".xlsb":
        return _iter_xlsb_rows(raw_bytes)
    if ext == ".xls":
        return _iter_xls_rows(raw_bytes)
    raise TabularLoaderError(f"Formato tabular não suportado: {ext or 'desconhecido'}")


def _iter_delimited_rows(raw_bytes: bytes, delimiter_hint: str | None) -> tuple[list[str], Iterator[list[str]]]:
    from tabular_loader import _sniff_delimiter

    text = raw_bytes.decode("utf-8-sig", errors="replace")
    if not text.strip():
        raise TabularLoaderError("CSV/TSV vazio.")
    delimiter = delimiter_hint or _sniff_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    header = next(reader, None)
    if not header:
        raise TabularLoaderError("CSV/TSV sem header válido.")
    columns = _normalize_header_row(header)

    def _rows() -> Iterator[list[str]]:
        for raw_row in reader:
            yield _row_values_from_sequence(raw_row)

    return columns, _rows()


def _iter_xlsx_rows(raw_bytes: bytes) -> tuple[list[str], Iterator[list[str]]]:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    worksheet = workbook.active
    row_iter = worksheet.iter_rows(values_only=True)
    header = next(row_iter, None)
    if not header:
        workbook.close()
        raise TabularLoaderError("Excel vazio.")
    columns = _normalize_header_row(header)

    def _rows() -> Iterator[list[str]]:
        try:
            for row in row_iter:
                yield _row_values_from_sequence(row)
        finally:
            workbook.close()

    return columns, _rows()


def _iter_xlsb_rows(raw_bytes: bytes) -> tuple[list[str], Iterator[list[str]]]:
    from pyxlsb import open_workbook

    temp_ctx = _temporary_tabular_file(raw_bytes, ".xlsb")
    temp_path = temp_ctx.__enter__()
    workbook = open_workbook(temp_path)
    worksheet = workbook.get_sheet(1)
    row_iter = worksheet.rows()
    header = next(row_iter, None)
    if not header:
        workbook.close()
        temp_ctx.__exit__(None, None, None)
        raise TabularLoaderError("XLSB vazio.")
    columns = _normalize_header_row(cell.v for cell in header)

    def _rows() -> Iterator[list[str]]:
        try:
            for row in row_iter:
                yield _row_values_from_sequence(cell.v for cell in row)
        finally:
            workbook.close()
            temp_ctx.__exit__(None, None, None)

    return columns, _rows()


def _iter_xls_rows(raw_bytes: bytes) -> tuple[list[str], Iterator[list[str]]]:
    try:
        import pandas as pd
    except Exception as exc:  # pragma: no cover
        raise TabularLoaderError("Leitura de .xls requer pandas/xlrd no servidor.") from exc

    temp_ctx = _temporary_tabular_file(raw_bytes, ".xls")
    temp_path = temp_ctx.__enter__()
    try:
        frame = pd.read_excel(temp_path, dtype=object)
    except Exception as exc:
        temp_ctx.__exit__(None, None, None)
        raise TabularLoaderError("Falha a ler ficheiro .xls.") from exc
    temp_ctx.__exit__(None, None, None)
    if frame.empty and not list(frame.columns):
        raise TabularLoaderError("Excel vazio.")
    columns = _normalize_header_row(frame.columns.tolist())

    def _rows() -> Iterator[list[str]]:
        for row in frame.itertuples(index=False, name=None):
            yield _row_values_from_sequence(row)

    return columns, _rows()
